# Variables dictionary for easy cycling
variables = {
    "#CharacterName1": "Alex",
    "#CharacterName1short": "Al",
    "#Item1": "magic sword",
    "#Item2": "golden key",
    "#Item3": "healing potion",
    "#CharacterName2": "Luna",
    "#Emote1": "smiled warmly",
    "#Gazed": "the ancient temple",
    "#Item4": "crystal orb",
    "#Item5": "leather boots",
    "#CharacterName3": "Marcus",
    "#Tøj1": "blue cloak",
    "#VoiceLine1": "Let's go on an adventure!",
    "#Hastighed": "lightning fast",
    "#SmukkereEndHvem": "a sunset",
    "#Item6": "enchanted ring",
    "#Placement": "on the wooden table",
    "#KeybindToWhat": "cast fireball",
    "#Keybind1": "Ctrl+F",
    "#Mad1": "hearty stew",
    "#Boss1": "Dragon King",
    "#Music1": "Epic Adventure",
    "#Music1Lyrics": "Rising up to meet our fate",
    "#Chair": "ornate throne",
}

# import time
import time

# Import tkinter
import tkinter as tk
root = tk.Tk()

# Import story
from docx import Document
doc = Document("story.docx")


def makeStory():
    # Prompt user for each variable
    for placeholder in variables.keys():
        user_input = input(f"Enter {placeholder} (press Enter to use '{variables[placeholder]}'): ")
        if user_input.strip():  # If user entered something, use their input
            variables[placeholder] = user_input
    
    # Replace placeholders in the document
    for parg in doc.paragraphs:
        for placeholder, value in variables.items():
            parg.text = parg.text.replace(placeholder, value + " ")
    

# Call the function to start the story generation
makeStory()

doc.save("SigmaStory.docx")
print("Story saved as SigmaStory.docx! (Replaced any existing file)")
